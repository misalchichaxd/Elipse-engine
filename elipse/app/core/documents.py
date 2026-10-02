"""
app/core/documents.py — Generación de documentos (Markdown -> PDF / Word).

Un solo parser de Markdown (una sola pasada por documento) alimenta a los dos
renderers, así el PDF y el Word salen iguales. Todo se guarda en
workspace/documents/ (dentro del sandbox).

  - Solo se entregan los formatos pedidos: pedir "pdf" produce UN archivo.
  - La fuente Markdown vive en documents/_src/ (uso interno para mode='append');
    no se lista ni se muestra al usuario, salvo que pida format='md'.
  - Las imágenes se descargan en paralelo, con tope de tamaño y validando cada
    redirección (anti-SSRF). Los fallos no se cachean.
  - Cada archivo se escribe a un .tmp y se mueve al final: nunca queda un PDF
    o DOCX a medias que el verificador de tools tome por bueno.

Dependencias: python-docx, reportlab (puro Python, andan en Windows).
"""
import io
import ipaddress
import os
import re
import socket
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import urlparse
from xml.sax.saxutils import escape

IMG_W_LAND, IMG_W_PORT = 4.0, 3.0    # ancho de las imágenes en pulgadas (apaisadas / verticales)
DOCS_SUBDIR = "documents"
SRC_SUBDIR = "_src"
FORMATS = ("pdf", "docx", "both", "md")


# -------------------------------------------------------------- imágenes ----

MAX_IMG_BYTES = 6_000_000
MAX_REDIRECTS = 3
_IMG_CACHE_MAX = 64
_img_cache: dict = {}
_img_lock = threading.Lock()


def _public_url(url: str) -> bool:
    p = urlparse(url)
    if p.scheme not in ("http", "https") or not p.hostname:
        return False
    try:
        for info in socket.getaddrinfo(p.hostname, None):
            ip = ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
                return False
    except Exception:
        return False
    return True


def _download(url: str):
    """Descarga con tope de bytes y revalidando cada redirección. -> bytes | None."""
    import httpx
    with httpx.Client(timeout=12, follow_redirects=False,
                      headers={"User-Agent": "Mozilla/5.0 (ELIPSE)"}) as client:
        for _ in range(MAX_REDIRECTS + 1):
            if not _public_url(url):
                return None
            with client.stream("GET", url) as r:
                if r.is_redirect:
                    loc = r.headers.get("location")
                    if not loc:
                        return None
                    url = str(httpx.URL(url).join(loc))
                    continue
                r.raise_for_status()
                buf = bytearray()
                for chunk in r.iter_bytes():
                    buf += chunk
                    if len(buf) > MAX_IMG_BYTES:
                        return None
                return bytes(buf)
    return None


def fetch_image(url: str):
    """Descarga una imagen pública y la normaliza a JPEG. -> (bytes, ancho, alto) o None."""
    with _img_lock:
        hit = _img_cache.get(url)
    if hit:
        return hit
    result = None
    try:
        raw = _download(url)
        if raw:
            from PIL import Image
            im = Image.open(io.BytesIO(raw))
            if im.mode in ("RGBA", "LA", "P"):        # transparencias sobre fondo blanco
                im = im.convert("RGBA")
                bg = Image.new("RGB", im.size, (255, 255, 255))
                bg.paste(im, mask=im.getchannel("A"))
                im = bg
            else:
                im = im.convert("RGB")
            im.thumbnail((1400, 1400))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=85)
            result = (buf.getvalue(), im.width, im.height)
    except Exception:
        result = None
    if result:                                        # los fallos NO se cachean (pueden ser transitorios)
        with _img_lock:
            if len(_img_cache) >= _IMG_CACHE_MAX:
                _img_cache.pop(next(iter(_img_cache)), None)
            _img_cache[url] = result
    return result


def _prefetch_images(blocks):
    """Baja todas las imágenes del documento en paralelo antes de renderizar."""
    urls = list({b[2] for b in blocks if b[0] == "img"})
    if len(urls) > 1:
        with ThreadPoolExecutor(max_workers=min(6, len(urls))) as ex:
            list(ex.map(fetch_image, urls))


# ---------------------------------------------------------------- parser ----

_INLINE = re.compile(
    r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\n]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)
_LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
_RE_IMG = re.compile(r"^!\[(.*)\]\(([^)\s]+)\)$")   # el alt puede traer [corchetes]
_RE_H = re.compile(r"^#{1,4}\s")
_RE_HR = re.compile(r"^(-{3,}|\*{3,}|_{3,})$")
_RE_TABLE_SEP = re.compile(r"^\s*\|?[\s:\-|]+\|?\s*$")
_RE_UL = re.compile(r"^\s*[-*+]\s+")
_RE_OL = re.compile(r"^\s*\d+[.)]\s+")
_RE_PARA_STOP = re.compile(r"^(#{1,4}\s|```|>|\||!\[|[-*+]\s|\d+[.)]\s|-{3,}$)")
_RE_TITLE = re.compile(r"^#\s+(.+)$", re.M)
_RE_OUTLINE = re.compile(r"^#{1,3}\s+(.+)$", re.M)


def parse_inline(text: str):
    """-> lista de (texto, estilo, url). estilo: '', 'b', 'i', 'c', 'a'."""
    out = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith(("**", "__")) and part.endswith(("**", "__")) and len(part) > 4:
            out.append((part[2:-2], "b", None))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append((part[1:-1], "c", None))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], "i", None))
        elif part.startswith("["):
            m = _LINK.match(part)
            out.append((m.group(1), "a", m.group(2)) if m else (part, "", None))
        else:
            out.append((part, "", None))
    return out


def parse_markdown(md: str):
    """-> bloques: ('h',nivel,txt) ('p',txt) ('ul',[items]) ('ol',[items])
    ('code',txt) ('quote',txt) ('table',[filas]) ('img',alt,url) ('hr',)"""
    lines = md.replace("\r\n", "\n").split("\n")
    n = len(lines)
    blocks, i = [], 0
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1
        elif s.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(buf)))
        elif (m := _RE_IMG.match(s)):
            blocks.append(("img", m.group(1), m.group(2)))
            i += 1
        elif _RE_H.match(s):
            lvl = len(s) - len(s.lstrip("#"))
            blocks.append(("h", lvl, s[lvl:].strip()))
            i += 1
        elif _RE_HR.match(s):
            blocks.append(("hr",))
            i += 1
        elif s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(buf)))
        elif s.startswith("|") and i + 1 < n and _RE_TABLE_SEP.match(lines[i + 1]) and "-" in lines[i + 1]:
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                if not _RE_TABLE_SEP.match(lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", rows))
        elif re.match(r"^[-*+]\s+", s):
            items = []
            while i < n and _RE_UL.match(lines[i]):
                items.append(_RE_UL.sub("", lines[i], count=1))
                i += 1
            blocks.append(("ul", items))
        elif re.match(r"^\d+[.)]\s+", s):
            items = []
            while i < n and _RE_OL.match(lines[i]):
                items.append(_RE_OL.sub("", lines[i], count=1))
                i += 1
            blocks.append(("ol", items))
        else:
            buf = []
            while i < n and lines[i].strip() and not _RE_PARA_STOP.match(lines[i].strip()):
                buf.append(lines[i].strip())
                i += 1
            if not buf:  # seguro anti-loop
                buf, i = [lines[i].strip()], i + 1
            blocks.append(("p", " ".join(buf)))
    return blocks

# ------------------------------------------------------------------ DOCX ----
#
# Formato formal (todo en negro, Arial):
#   Título del documento y encabezados nivel 1-2 ...... Arial 14 negrita
#   Subtítulos (encabezados nivel 3-4) ................ Arial 13 negrita
#   Texto normal ...................................... Arial 12, interlineado 1.5, justificado

FONT_NAME = "Arial"
SIZE_TITLE = 14      # título del documento y secciones (# y ##)
SIZE_SUBTITLE = 13   # subtítulos (### y ####)
SIZE_BODY = 12
SIZE_TABLE = 11
LINE_SPACING = 1.5
ARIAL_LINE = 1.15    # altura de línea de Arial (Word calcula el interlineado sobre esto)


def build_docx(md: str, path: Path, title: str = "", author: str = "ELIPSE", blocks=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    BLACK = RGBColor(0, 0, 0)

    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = author
    sec = doc.sections[0]                      # Carta con márgenes de 2.54 cm, igual que el PDF
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)

    def style_font(style, size, bold=None):
        """Arial + tamaño + negro. Quita los atributos de tema (si no, Word ignora la fuente)."""
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.italic = False
        if bold is not None:
            style.font.bold = bold
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for att in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rfonts.attrib.pop(qn(f"w:{att}"), None)
        for att in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{att}"), FONT_NAME)

    # Texto normal (las listas heredan de aquí)
    normal = doc.styles["Normal"]
    style_font(normal, SIZE_BODY)
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(6)

    # Título del documento (con su línea inferior en negro)
    t = doc.styles["Title"]
    style_font(t, SIZE_TITLE, bold=True)
    t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.line_spacing = LINE_SPACING
    t.paragraph_format.space_after = Pt(12)
    for el in t.element.iter(qn("w:bottom")):
        el.set(qn("w:color"), "000000")
        el.set(qn("w:val"), "single")
        for att in ("themeColor", "themeShade", "themeTint"):   # el tema pisaría el negro
            el.attrib.pop(qn(f"w:{att}"), None)

    # Encabezados
    for name, size in (("Heading 1", SIZE_TITLE), ("Heading 2", SIZE_TITLE),
                       ("Heading 3", SIZE_SUBTITLE), ("Heading 4", SIZE_SUBTITLE)):
        hs = doc.styles[name]
        style_font(hs, size, bold=True)
        hs.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hs.paragraph_format.line_spacing = LINE_SPACING
        hs.paragraph_format.space_before = Pt(12)
        hs.paragraph_format.space_after = Pt(6)
        hs.paragraph_format.keep_with_next = True

    def add_runs(par, text, size=None):
        for txt, sty, url in parse_inline(text):
            r = par.add_run(txt)
            r.bold = sty == "b" or None
            r.italic = sty == "i" or None
            r.font.color.rgb = BLACK
            if size:
                r.font.size = Pt(size)
            if sty == "c":
                r.font.name, r.font.size = "Consolas", Pt(10)
            if sty == "a":
                r.underline = True
                if url:
                    u = par.add_run(f" ({url})")
                    u.font.size = Pt(10)
                    u.font.color.rgb = BLACK

    def plain(par, align=None, spacing=1.0, after=0):
        par.paragraph_format.line_spacing = spacing
        par.paragraph_format.space_after = Pt(after)
        if align is not None:
            par.alignment = align

    if title:
        doc.add_heading(title, level=0)

    for b in (blocks if blocks is not None else parse_markdown(md)):
        k = b[0]
        if k == "h":
            doc.add_heading(b[2], level=min(b[1], 4))
        elif k == "p":
            add_runs(doc.add_paragraph(), b[1])
        elif k in ("ul", "ol"):
            style = "List Bullet" if k == "ul" else "List Number"
            for it in b[1]:
                add_runs(doc.add_paragraph(style=style), it)
        elif k == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            add_runs(p, b[1])
            for r in p.runs:
                r.italic = True
        elif k == "code":
            p = doc.add_paragraph()
            plain(p, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 8)
            r = p.add_run(b[1])
            r.font.name, r.font.size = "Consolas", Pt(9)
            r.font.color.rgb = BLACK
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F0EFEC")
            p._p.get_or_add_pPr().append(shd)
        elif k == "img":
            got = fetch_image(b[2])
            p = doc.add_paragraph()
            plain(p, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 4)
            if got:
                p.add_run().add_picture(io.BytesIO(got[0]), width=Inches(IMG_W_LAND) if got[1] >= got[2] else Inches(IMG_W_PORT))
                if b[1]:
                    c = doc.add_paragraph()
                    plain(c, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 8)
                    r = c.add_run(b[1])
                    r.italic = True
                    r.font.size = Pt(10)
                    r.font.color.rgb = BLACK
            else:
                p.add_run(f"[imagen no disponible: {b[1] or b[2]}]").italic = True
        elif k == "hr":
            p = doc.add_paragraph()
            plain(p, WD_ALIGN_PARAGRAPH.CENTER, 1.0, 6)
            p.add_run("─" * 30).font.color.rgb = BLACK
        elif k == "table" and b[1]:
            rows = b[1]
            ncol = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=ncol)
            tbl.style = "Table Grid"          # bordes negros simples
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = tbl.cell(ri, ci)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    plain(par, WD_ALIGN_PARAGRAPH.LEFT, 1.0, 0)
                    add_runs(par, row[ci] if ci < len(row) else "", size=SIZE_TABLE)
                    if ri == 0:
                        for r in par.runs:
                            r.bold = True
            doc.add_paragraph()
    doc.save(str(path))


# ------------------------------------------------------------------- PDF ----

_PDF_FONT = None


def _pdf_fonts():
    """Registra Arial real (TTF de Windows) si existe; si no, Helvetica (equivalente métrico).
    -> (regular, negrita, cursiva)"""
    global _PDF_FONT
    if _PDF_FONT:
        return _PDF_FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    dirs = [Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts",
            Path("/usr/share/fonts/truetype/msttcorefonts"),
            Path("/Library/Fonts"), Path("/System/Library/Fonts/Supplemental")]
    files = ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf")
    for d in dirs:
        if all((d / f).exists() for f in files):
            try:
                for name, f in zip(("Arial", "Arial-Bold", "Arial-Italic", "Arial-BoldItalic"), files):
                    pdfmetrics.registerFont(TTFont(name, str(d / f)))
                pdfmetrics.registerFontFamily("Arial", normal="Arial", bold="Arial-Bold",
                                              italic="Arial-Italic", boldItalic="Arial-BoldItalic")
                _PDF_FONT = ("Arial", "Arial-Bold", "Arial-Italic")
                return _PDF_FONT
            except Exception:
                pass
    _PDF_FONT = ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique")
    return _PDF_FONT


def _rl_inline(text: str) -> str:
    out = []
    for t, sty, url in parse_inline(text):
        t = escape(t)
        if sty == "b":
            t = f"<b>{t}</b>"
        elif sty == "i":
            t = f"<i>{t}</i>"
        elif sty == "c":
            t = f'<font face="Courier" size="10">{t}</font>'
        elif sty == "a":
            t = f'<link href="{escape(url or "")}" color="#000000"><u>{t}</u></link>'
        out.append(t)
    return "".join(out)


def build_pdf(md: str, path, title: str = "", author: str = "ELIPSE", blocks=None, on_page=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (HRFlowable, KeepTogether, ListFlowable, ListItem,
                                    Paragraph, Preformatted, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)
    from reportlab.platypus import Image as RLImage

    F, FB, FI = _pdf_fonts()
    BLACK = colors.black

    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontName=F, fontSize=SIZE_BODY,
                          leading=SIZE_BODY * ARIAL_LINE * LINE_SPACING, alignment=TA_JUSTIFY,
                          textColor=BLACK, spaceAfter=6)
    sizes = {1: SIZE_TITLE, 2: SIZE_TITLE, 3: SIZE_SUBTITLE, 4: SIZE_SUBTITLE}
    h = {n: ParagraphStyle(f"h{n}", parent=body, fontName=FB, fontSize=s,
                           leading=s * ARIAL_LINE * LINE_SPACING, alignment=TA_LEFT,
                           spaceBefore=12, spaceAfter=6, textColor=BLACK, keepWithNext=1)
         for n, s in sizes.items()}
    quote = ParagraphStyle("q", parent=body, leftIndent=18, fontName=FI, textColor=BLACK)
    code = ParagraphStyle("c", fontName="Courier", fontSize=9, leading=12, textColor=BLACK,
                          backColor=colors.HexColor("#F0EFEC"), borderPadding=6, spaceAfter=8)
    cell = ParagraphStyle("cell", parent=body, fontSize=SIZE_TABLE, leading=SIZE_TABLE * ARIAL_LINE * 1.1,
                          alignment=TA_LEFT, spaceAfter=0)
    cell_head = ParagraphStyle("cellh", parent=cell, fontName=FB)

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont(F, 9)
        canvas.setFillColor(BLACK)
        if on_page:
            on_page(d.page)
        canvas.drawCentredString(letter[0] / 2, 1.3 * cm, f"{title}  ·  {d.page}" if title else str(d.page))
        canvas.restoreState()

    story = []
    if title:
        ts = ParagraphStyle("t", parent=h[1], alignment=TA_CENTER, spaceBefore=0, spaceAfter=8)
        story += [Paragraph(escape(title), ts),
                  HRFlowable(width="100%", thickness=1, color=BLACK), Spacer(1, 12)]

    for b in (blocks if blocks is not None else parse_markdown(md)):
        k = b[0]
        if k == "h":
            story.append(Paragraph(_rl_inline(b[2]), h[min(b[1], 4)]))
        elif k == "p":
            story.append(Paragraph(_rl_inline(b[1]), body))
        elif k in ("ul", "ol"):
            items = [ListItem(Paragraph(_rl_inline(x), body)) for x in b[1]]
            story.append(ListFlowable(items, bulletType="bullet" if k == "ul" else "1",
                                      leftIndent=18, bulletFontName=F, bulletFontSize=SIZE_BODY))
        elif k == "quote":
            story.append(Paragraph(_rl_inline(b[1]), quote))
        elif k == "code":
            story.append(Preformatted(b[1], code))
        elif k == "img":
            got = fetch_image(b[2])
            if got:
                w_in = IMG_W_LAND if got[1] >= got[2] else IMG_W_PORT   # mismo tamaño que en el Word
                sc = (w_in * 2.54 * cm) / got[1]
                img = RLImage(io.BytesIO(got[0]), width=got[1] * sc, height=got[2] * sc)
                cap = ParagraphStyle("cap", parent=body, alignment=TA_CENTER, fontSize=10,
                                     leading=13, fontName=FI, textColor=BLACK)
                story.append(KeepTogether([img, Paragraph(_rl_inline(b[1]), cap) if b[1] else Spacer(1, 4)]))
            else:
                story.append(Paragraph(f"<i>[imagen no disponible: {escape(b[1] or b[2])}]</i>", body))
        elif k == "hr":
            story.append(HRFlowable(width="100%", thickness=0.8, color=BLACK,
                                    spaceBefore=6, spaceAfter=6))
        elif k == "table" and b[1]:
            rows = b[1]
            n = max(len(r) for r in rows)
            data = [[Paragraph(_rl_inline(r[c]) if c < len(r) else "", cell_head if ri == 0 else cell)
                     for c in range(n)] for ri, r in enumerate(rows)]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.6, BLACK),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]))
            story += [t, Spacer(1, 8)]

    target = path if hasattr(path, "write") else str(path)
    SimpleDocTemplate(target, pagesize=letter, title=title, author=author,
                      leftMargin=2.54 * cm, rightMargin=2.54 * cm,
                      topMargin=2.54 * cm, bottomMargin=2.54 * cm
                      ).build(story, onFirstPage=footer, onLaterPages=footer)
def md_safe_url(url: str) -> str:
    return url.replace(" ", "%20").replace("(", "%28").replace(")", "%29")

def estimate_pages(text: str) -> int:
    """Páginas que ocupará el documento (se mide armando el PDF en memoria; el Word queda casi igual,
    porque usa la misma hoja, márgenes, fuente e interlineado)."""
    m = _RE_TITLE.search(text)
    title = m.group(1).strip() if m else ""
    body = re.sub(r"^#\s+.+\n", "", text, count=1) if re.match(r"^#\s+", text) else text
    blocks = parse_markdown(body)
    _prefetch_images(blocks)
    seen = {"n": 1}
    build_pdf(body, io.BytesIO(), title, blocks=blocks, on_page=lambda n: seen.__setitem__("n", max(seen["n"], n)))
    return seen["n"]


# ------------------------------------------------------------ orquestador ----

def _slug(name: str) -> str:
    base = Path(name).stem
    base = re.sub(r"[^\w\-]+", "_", base, flags=re.UNICODE).strip("_")
    return (base or "documento")[:80]


def _atomic(path: Path, write):
    """write(tmp_path) y luego mueve a 'path': nunca queda un archivo a medias."""
    tmp = path.with_name(path.name + ".tmp")
    try:
        write(tmp)
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def generate_document(root: Path, filename: str, content: str, fmt: str = "pdf",
                      title: str = "", mode: str = "create") -> dict:
    """
    root: WORKSPACE_ROOT. Renderiza SOLO los formatos pedidos (pdf, docx, both, md).
    La fuente Markdown queda en documents/_src/<slug>.md (interna, para mode='append');
    solo se publica un .md en documents/ si fmt == 'md'.
    Devuelve {status, title, files, chars, outline} o {status:'error', detail}.
    """
    fmt = (fmt or "pdf").lower()
    if fmt not in FORMATS:
        return {"status": "error", "detail": "formato debe ser pdf, docx, both o md."}
    if not content or not content.strip():
        return {"status": "error", "detail": "el contenido está vacío."}

    folder = root / DOCS_SUBDIR
    src_dir = folder / SRC_SUBDIR
    src_dir.mkdir(parents=True, exist_ok=True)
    slug = _slug(filename)
    src = src_dir / f"{slug}.md"

    if mode == "append" and src.exists():
        text = src.read_text(encoding="utf-8").rstrip() + "\n\n" + content.strip() + "\n"
    else:
        text = content.strip() + "\n"
    _atomic(src, lambda p: p.write_text(text, encoding="utf-8"))

    m = _RE_TITLE.search(text)
    title = title.strip() or (m.group(1).strip() if m else slug.replace("_", " "))
    # si el md ya abre con "# titulo", no repetirlo: el título va como portada
    body_md = re.sub(r"^#\s+.+\n", "", text, count=1) if re.match(r"^#\s+", text) else text

    files = []
    try:
        if fmt == "md":
            out = folder / f"{slug}.md"
            _atomic(out, lambda p: p.write_text(text, encoding="utf-8"))
            files.append(f"{DOCS_SUBDIR}/{out.name}")
        if fmt in ("pdf", "docx", "both"):
            blocks = parse_markdown(body_md)          # una sola pasada para ambos renderers
            _prefetch_images(blocks)
            if fmt in ("pdf", "both"):
                out = folder / f"{slug}.pdf"
                _atomic(out, lambda p: build_pdf(body_md, p, title, blocks=blocks))
                files.append(f"{DOCS_SUBDIR}/{out.name}")
            if fmt in ("docx", "both"):
                out = folder / f"{slug}.docx"
                _atomic(out, lambda p: build_docx(body_md, p, title, blocks=blocks))
                files.append(f"{DOCS_SUBDIR}/{out.name}")
    except Exception as e:
        return {"status": "error", "detail": f"falló el render: {e}"}

    return {"status": "ok", "title": title, "files": files, "chars": len(text),
            "outline": _RE_OUTLINE.findall(text)}